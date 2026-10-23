#!/usr/bin/env python3
"""特許明細書の事務校正（ルール + ローカルLLM）

使い方:
    python proofread.py 明細書.docx
    python proofread.py 明細書.docx --no-llm          # ルールチェックだけ
    python proofread.py 明細書.docx --base-url http://xxxx/v1 --model モデル名

出力（入力と同じフォルダ）:
    【AI校正】明細書.docx      … 変更履歴（明確な誤記）とコメント（要確認・申し送り）付き
    明細書_修正箇所一覧.csv    … 全指摘の一覧
    明細書_符号対応表.csv      … 名称と符号の対応表（矛盾の種類つき）

依存: python-docx、janome（語の重複チェック。なければ省略）、pypdf（図面照合）
"""
from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

DEFAULT_BASE_URL = os.environ.get("PROOF_LLM_BASE_URL", "http://3a77b15d.shared.olares.com/v1")
DEFAULT_MODEL = os.environ.get("PROOF_LLM_MODEL", "Olares/unsloth/Qwen3.8-27B-GGUF:UD-Q4_K_XL")
AUTHOR = "AI校正"
REASONING = os.environ.get("PROOF_LLM_REASONING", "low")  # low / medium / xhigh / 空文字=モデル既定

# ---------------------------------------------------------------------------
# 校正方針（既定）。事務所ごとの校正マニュアルがあれば、アプリの設定画面や
# 環境変数 PROOF_MANUAL_FILE（テキストファイル）で差し替えられる
# ---------------------------------------------------------------------------
DEFAULT_MANUAL = """\
（特許明細書の事務校正。技術内容ではなく、書誌的な誤記を確認する）
- 確認する: 誤字・脱字・誤変換、助詞の誤り、用語（名詞）の不統一、漢字／ひらがな／カタカナの表記ゆれ、口語的・明らかに不自然な表現、符号と名称の対応の食い違い、図番号の誤り。
- 確認しない: 技術内容の当否、法的な観点、文体の好み、事務所や担当者固有の書き方ルール、願書部分。
- 明確な誤記は修正案を示す。修正してよいか迷う箇所は「要確認」として申し送る。
- 違和感があっても誤りだと説明できない箇所は指摘しない。"""


def _load_manual() -> str:
    path = os.environ.get("PROOF_MANUAL_FILE")
    if path and os.path.exists(path):
        return open(path, encoding="utf-8").read().strip() or DEFAULT_MANUAL
    return DEFAULT_MANUAL


MANUAL = _load_manual()

TYPES = ("タイポ", "不自然な表現", "用語の不統一", "図面・符号との一致")

FW_DIGIT = "０-９"
FW_UPPER = "Ａ-Ｚ"


# ---------------------------------------------------------------------------
# データ構造
# ---------------------------------------------------------------------------
@dataclass
class Para:
    idx: int          # document.paragraphs のインデックス
    label: str        # 【００２４】 など
    text: str


@dataclass
class Finding:
    source: str       # rule / llm-whole / llm-seg
    label: str
    before: str
    after: str
    kind: str
    note: str
    para_idx: int | None = None
    start: int | None = None
    needs_check: bool = False
    applied: str = ""  # change / comment / report-only
    votes: int = 0    # LLM 多数決の検出回数（0 = ルールによる指摘）

    def key(self) -> str:
        return re.sub(r"\s", "", self.before)


# ---------------------------------------------------------------------------
# 1. docx → 段落（願書部分は除外）
# ---------------------------------------------------------------------------
LABEL_RE = re.compile(r"^【([^】]+)】")


def is_para_label(inner: str) -> bool:
    return bool(re.fullmatch(f"[{FW_DIGIT}]+", inner) or inner.startswith("請求項") or inner.startswith("図") or inner == "符号の説明")


def extract_paragraphs(doc) -> list[Para]:
    paras: list[Para] = []
    label = ""
    in_body = False
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if re.match(r"^【書類名】\s*明細書", t):
            in_body = True
        if not in_body:
            continue
        m = LABEL_RE.match(t)
        if m and is_para_label(m.group(1)):
            label = "【" + m.group(1) + "】"
        elif m and not label:
            label = "【" + m.group(1) + "】"
        paras.append(Para(i, label, t))
    return paras


# ---------------------------------------------------------------------------
# 2. ルールチェック（名称と符号の対応表・語の重複・表記ゆれ）
# ---------------------------------------------------------------------------
NUM_RE = re.compile(rf"([ァ-ヶー一-龠々]{{2,}})([{FW_DIGIT}]{{1,4}}[{FW_UPPER}]?)(?![{FW_DIGIT}．{FW_UPPER}])")
PREFIX_RE = re.compile(r"^(前記|上記|当該|該|各|及び|又は|並びに|若しくは)+")
# 符号ではない「語＋数字」（文献番号・請求項・図番号・日付・式番号など）
EXCLUDE_TERM = re.compile(r"(請求項|特許第|第|図|段落|ステップ|文献|番号|工程|特開|特表|再表|特願|実開|実願|国際公開|公報|式|数|表|頁|年|月|日)$")
# 【符号の説明】の1項目（例: １０…撹拌装置 / １０　撹拌装置 / 10：撹拌装置）
LEGEND_RE = re.compile(rf"(?<![{FW_DIGIT}0-9])([{FW_DIGIT}0-9]{{1,4}}[{FW_UPPER}A-Z]?)\s*(?:…+|‥+|：|:|－|-|\s)+\s*([^\s、，,…‥：:{FW_DIGIT}0-9][^\s、，,…‥：:]*)")
Z2H = str.maketrans("０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ", "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ")


def ctx(text: str, start: int, end: int, pad: int = 8) -> str:
    return text[max(0, start - pad): end + pad]


def is_legend(p: Para) -> bool:
    return p.label == "【符号の説明】" or "…" in p.text


def symbol_occurrences(paras: list[Para]) -> list[tuple]:
    """本文中の「名称＋符号」を拾う。(名称, 符号(全角のまま), 段落, 名称の開始位置, 終了位置)"""
    occ = []
    for p in paras:
        if is_legend(p):
            continue
        for m in NUM_RE.finditer(p.text):
            term = PREFIX_RE.sub("", m.group(1))
            if len(term) < 2 or EXCLUDE_TERM.search(term):
                continue
            occ.append((term, m.group(2), p, m.start(1) + (len(m.group(1)) - len(term)), m.end()))
    return occ


def legend_entries(paras: list[Para]) -> dict[str, tuple]:
    """【符号の説明】の 符号(半角) → (名称, 段落, 位置, 符号の表記)"""
    out: dict[str, tuple] = {}
    for p in paras:
        if not is_legend(p):
            continue
        for m in LEGEND_RE.finditer(p.text):
            name = PREFIX_RE.sub("", m.group(2))
            if len(name) >= 2:
                out.setdefault(m.group(1).translate(Z2H), (name, p, m.start(2), m.group(1)))
    return out


def _related(a: str, b: str) -> bool:
    """同じものを指すとみなす名称（一方が他方の末尾・先頭に含まれる。例: 羽根 / 撹拌羽根）"""
    return a == b or a.endswith(b) or b.endswith(a) or a.startswith(b) or b.startswith(a)


def _digit_slip(a: str, b: str) -> bool:
    """符号の打ち間違いらしい差（数字の脱落・重複・入れ替わり）。末尾1桁だけの違いは兄弟の符号とみなして除く"""
    if abs(len(a) - len(b)) == 1:
        long, short = (a, b) if len(a) > len(b) else (b, a)
        return any(long[:i] + long[i + 1:] == short for i in range(len(long)))
    if len(a) == len(b) and len(a) >= 2:
        d = [i for i in range(len(a)) if a[i] != b[i]]
        return len(d) == 2 and d[1] == d[0] + 1 and a[d[0]] == b[d[1]] and a[d[1]] == b[d[0]]
    return False


def symbol_table(paras: list[Para]) -> dict[str, dict]:
    """名称と符号の対応表。符号(半角) → {fw: 全角表記, names: {名称: 回数}, occ: [...], legend: (名称, 段落, 位置) or None}"""
    table: dict[str, dict] = {}
    for o in symbol_occurrences(paras):
        e = table.setdefault(o[1].translate(Z2H), {"fw": o[1], "names": {}, "occ": [], "legend": None, "fig": None, "issues": []})
        e["names"][o[0]] = e["names"].get(o[0], 0) + 1
        e["occ"].append(o)
    for num, ent in legend_entries(paras).items():
        table.setdefault(num, {"fw": ent[3], "names": {}, "occ": [], "legend": None, "fig": None, "issues": []})["legend"] = ent
    return table


def symbol_checks(paras: list[Para], table: dict[str, dict]) -> list[Finding]:
    """対応表の矛盾を探す。
    R1a 同じ符号に別の名称 / R1b 1回しか出ない符号 / R1c 同じ名称で数字の打ち間違いらしい符号 /
    R1d 【符号の説明】と本文の名称の食い違い / R1e 【符号の説明】と本文の過不足"""
    out: list[Finding] = []

    def add(ent, o, after, kind, note, tag):
        term, fw, p, s, e = o
        ent["issues"].append(tag)
        out.append(Finding("rule", p.label, term + fw, after, kind, note, p.idx, s, needs_check=True))

    has_legend = sum(1 for e in table.values() if e["legend"]) >= 3
    by_name: dict[str, dict[str, int]] = {}
    for num, ent in table.items():
        for term, n in ent["names"].items():
            by_name.setdefault(term, {})[num] = n

    for num, ent in table.items():
        cnt, lst = ent["names"], ent["occ"]
        if cnt:
            # R1a
            main = max(cnt, key=cnt.get)
            for term in cnt:
                if not _related(term, main):
                    o = next(x for x in lst if x[0] == term)
                    add(ent, o, main + ent["fw"], "用語の不統一", f"符号{ent['fw']}は他の箇所では「{main}」（{cnt[main]}回）。", "名称の食い違い")
            # R1c
            for term, n in cnt.items():
                if n > 2:
                    continue
                syms = by_name.get(term, {})
                if len(syms) > 3:  # 同じ名称に符号が多い（流路１３０１〜１３０６ など）は連番とみなす
                    continue
                for other, m in syms.items():
                    if other != num and m >= 3 and _digit_slip(num, other) and not (ent["legend"] and _related(ent["legend"][0], term)):
                        o = next(x for x in lst if x[0] == term)
                        add(ent, o, term + table[other]["fw"], "図面・符号との一致",
                            f"「{term}」の符号は他の箇所では「{table[other]['fw']}」（{m}回）。数字の打ち間違いの可能性。", "符号の打ち間違い？")
                        break
            # R1d
            if ent["legend"]:
                lname = ent["legend"][0]
                for term in cnt:
                    if not _related(term, lname):
                        o = next(x for x in lst if x[0] == term)
                        add(ent, o, lname + ent["fw"], "図面・符号との一致", f"【符号の説明】では符号{ent['fw']}は「{lname}」。", "符号の説明と不一致")
            # R1b / R1e（本文にあって符号の説明にない）。上で矛盾を指摘済みの符号には重ねない
            if ent["issues"]:
                pass
            elif len(lst) == 1 and not ent["legend"]:
                add(ent, lst[0], "", "図面・符号との一致",
                    f"符号「{lst[0][0]}{ent['fw']}」は文書中この1回しか出てこない。他に説明のない符号の可能性。", "1回だけ")
            elif has_legend and not ent["legend"]:
                add(ent, lst[0], "", "図面・符号との一致", f"符号「{lst[0][0]}{ent['fw']}」が【符号の説明】にない。", "符号の説明にない")
        elif ent["legend"]:
            # R1e（符号の説明にあって本文にない）
            name, p, i, _ = ent["legend"]
            ent["issues"].append("本文に出てこない")
            out.append(Finding("rule", p.label, name, "", "図面・符号との一致",
                               f"【符号の説明】の符号{ent['fw']}「{name}」が本文に出てこない。", p.idx, i, needs_check=True))
    return out


def table_rows(table: dict[str, dict]) -> list[list[str]]:
    """対応表を CSV・画面用の行にする（符号の数字順）"""
    def key(k):
        m = re.match(r"(\d+)(.*)", k)
        return (int(m.group(1)), m.group(2)) if m else (10 ** 6, k)
    rows = []
    for num in sorted(table, key=key):
        e = table[num]
        names = "、".join(f"{t}（{n}）" for t, n in sorted(e["names"].items(), key=lambda x: -x[1]))
        first = e["occ"][0][2].label if e["occ"] else (e["legend"][1].label if e["legend"] else "")
        fig = {True: "○", False: "×", None: "—"}[e["fig"]]
        rows.append([e["fw"], names, e["legend"][0] if e["legend"] else "", first, fig, "、".join(dict.fromkeys(e["issues"]))])
    return rows


def rule_checks(paras: list[Para], table: dict[str, dict] | None = None) -> list[Finding]:
    if table is None:
        table = symbol_table(paras)
    return symbol_checks(paras, table) + hyoki_checks(paras) + repeat_checks(paras)


# 同じ語の重複（形態素解析）。例: 「のの」「装置装置」「受け取って受け取った」「貯留し貯留する」
_PUNCT_DUP = {"、", "。", "，", "．"}
_BETWEEN_OK = ("助動詞", "助詞,接続助詞")


def _pos(t) -> str:
    return t.part_of_speech


def morph_tokens(text: str, tk) -> list[tuple[str, str, str, int]]:
    """(表層, 品詞, 原形, 開始位置)"""
    out, pos = [], 0
    for t in tk.tokenize(text):
        i = text.find(t.surface, pos)
        i = pos if i < 0 else i
        out.append((t.surface, t.part_of_speech, t.base_form, i))
        pos = i + len(t.surface)
    return out


def _skip_unit(unit) -> bool:
    surf = "".join(u[0] for u in unit)
    if not surf.strip():
        return True
    if all(u[1].startswith("記号") for u in unit):
        return surf not in _PUNCT_DUP  # 「、、」「。。」だけを見る（記号の繰り返しは数式などで正当なことがある）
    if any(u[1].startswith("名詞,数") or u[1].startswith("記号,アルファベット") for u in unit):
        return True
    # 漢字・カタカナ1字の名詞の繰り返し（「人人」など）は除く。ひらがな1字（「のの」「をを」）は誤記として見る
    return len(unit) == 1 and len(surf) == 1 and unit[0][1].startswith("名詞") and not re.fullmatch(r"[ぁ-ん]", surf)


def repeat_checks(paras: list[Para]) -> list[Finding]:
    out: list[Finding] = []
    # 同じ文が続けて2回（正規表現。形態素解析が使えなくても動く）
    for p in paras:
        for m in re.finditer(r"([^。]{6,}。)\1", p.text):
            out.append(Finding("rule", p.label, m.group(0), m.group(1), "タイポ", "同じ文が続けて2回書かれている。", p.idx, m.start()))
    try:
        from janome.tokenizer import Tokenizer
    except ImportError:
        log("形態素解析（janome）が使えないため、語の重複チェックは同じ文の繰り返しだけ行います")
        return out
    tk = Tokenizer()
    for p in paras:
        toks = morph_tokens(p.text, tk)
        i = 0
        while i < len(toks):
            hit = None
            # A: 同じ語（1〜6語の並び）がそのまま続く
            for n in range(min(6, (len(toks) - i) // 2), 0, -1):
                unit, nxt = toks[i:i + n], toks[i + n:i + 2 * n]
                if [u[0] for u in unit] == [u[0] for u in nxt] and not _skip_unit(unit):
                    s, e = unit[0][3], nxt[-1][3] + len(nxt[-1][0])
                    word = p.text[s:unit[-1][3] + len(unit[-1][0])]
                    kana_word = bool(re.fullmatch(r"[ぁ-んァ-ヶー]{2,}", word)) and not unit[0][1].startswith(("助詞", "助動詞"))
                    a, b = max(0, s - 4), min(len(p.text), e + 4)
                    hit = Finding("rule", p.label, p.text[a:b], p.text[a:s] + word + p.text[e:b], "タイポ",
                                  f"語の重複（形態素解析）: 「{word}」が続けて2回出ている。" + (" 要確認" if kana_word else ""),
                                  p.idx, a, needs_check=kana_word)
                    i += 2 * n
                    break
            if hit:
                out.append(hit)
                continue
            # B: 同じ自立語が活用語尾・接続助詞だけを挟んで繰り返される（「貯留し貯留する」「受け取って受け取った」）
            t = toks[i]
            if (t[1].startswith("動詞,自立") or (t[1].startswith("名詞") and not t[1].startswith(("名詞,数", "名詞,非自立", "名詞,代名詞", "名詞,接尾")))) \
                    and len(t[0]) >= 2:
                for j in range(i + 1, min(i + 4, len(toks))):
                    u = toks[j]
                    if u[2] == t[2] and u[1].split(",")[0] == t[1].split(",")[0]:
                        between = toks[i + 1:j]
                        follow = toks[j + 1] if j + 1 < len(toks) else None
                        # 「記憶する記憶部」「取得する取得部」のように後ろが名詞につながる形は特許の定型表現なので除く
                        if between and all(b[1].startswith(_BETWEEN_OK) or b[2] == "する" for b in between) \
                                and not (follow and follow[1].startswith("名詞")):
                            s, e = t[3], u[3] + len(u[0])
                            a, b = max(0, s - 4), min(len(p.text), e + 6)
                            out.append(Finding("rule", p.label, p.text[a:b], "", "タイポ",
                                               f"語の重複（形態素解析）: 「{t[2]}」が続けて出ている。どちらかが不要の可能性。 要確認",
                                               p.idx, a, needs_check=True))
                            i = j
                        break
                    if not (u[1].startswith(_BETWEEN_OK) or u[2] == "する"):
                        break
            i += 1
    del tk
    return out


# カタカナ語の表記ゆれ（長音・小書き文字の違いだけの語。例: コンピュータ/コンピューター）
_SMALL = str.maketrans("ァィゥェォャュョ", "アイウエオヤユヨ")


def _kana_key(w: str) -> str:
    return w.rstrip("ー").replace("ー", "").translate(_SMALL)


def hyoki_checks(paras: list[Para]) -> list[Finding]:
    out: list[Finding] = []
    words: dict[str, dict[str, int]] = {}
    for p in paras:
        for m in re.finditer(r"[ァ-ヶー]{3,}", p.text):
            w = m.group(0)
            words.setdefault(_kana_key(w), {}).setdefault(w, 0)
            words[_kana_key(w)][w] += 1
    for key, forms in words.items():
        if len(forms) < 2:
            continue
        main = max(forms, key=forms.get)
        for w in forms:
            if w == main:
                continue
            hit = None
            for p in paras:
                for m in re.finditer(re.escape(w), p.text):
                    i = m.start()
                    # 長い語の一部として出ている場合は除く
                    if not re.match(r"[ァ-ヶー]", p.text[i + len(w):i + len(w) + 1] or "x") and \
                            not re.match(r"[ァ-ヶー]", p.text[i - 1:i] if i else "x"):
                        hit = (p, i)
                        break
                if hit:
                    break
            if hit:
                p, i = hit
                out.append(Finding("rule", p.label, w, main, "用語の不統一",
                                   f"表記ゆれ: 「{w}」と「{main}」（{forms[main]}回）が混在。 要確認", p.idx, i, needs_check=True))
    return out


OCR_URL = os.environ.get("PROOF_OCR_URL", "http://43279715.shared.olares.com")  # PaddleOCR の共有エントランス


def ocr_pdf_text(pdf_path) -> str | None:
    """PaddleOCR（layout-parsing API）で図面PDFの全ページの文字を読む。使えなければ None。"""
    if not OCR_URL:
        return None
    import base64
    data = base64.b64encode(open(pdf_path, "rb").read()).decode()
    body = json.dumps({"file": data, "fileType": 0, "useLayoutDetection": False,
                       "useChartRecognition": False, "visualize": False}).encode()
    for path in ("/layout-parsing", "/api/layout-parsing"):
        try:
            req = urllib.request.Request(OCR_URL.rstrip("/") + path, data=body, headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=900) as r:
                res = json.loads(r.read())
            pages = res["result"]["layoutParsingResults"]
            text = "\n".join(b.get("block_content", "") for pg in pages for b in pg["prunedResult"]["parsing_res_list"])
            log(f"図面PDF: OCR で {len(pages)} ページ / {len(text)} 文字を読み取り")
            return text
        except Exception as e:
            last = e
    log(f"図面PDF: OCR に接続できません（{OCR_URL}: {last}）")
    return None


def ocr_cache_path(pdf_path) -> Path:
    p = Path(pdf_path)
    return p.with_name(p.stem + "_OCR.txt")


def drawing_symbol_checks(paras: list[Para], pdf_path, allow_ocr: bool = True, table: dict | None = None) -> list[Finding]:
    """明細書に出てくる符号が図面PDFの文字にあるかを照合し、見つからない符号を申し送る。
    allow_ocr=True なら OCR で読み直す。False なら保存済みの OCR 結果を使う（なければ PDF 内の文字情報）。"""
    out: list[Finding] = []
    cache = ocr_cache_path(pdf_path)
    text = None
    if allow_ocr:
        # OCR は GPU を使うため、LLM（Qwen）と同時に動かさない。呼び出し側が LLM の工程より前に済ませる
        text = ocr_pdf_text(pdf_path)
        if text is not None:
            cache.write_text(text, encoding="utf-8")
            log(f"図面PDF: OCR 結果を {cache.name} に保存")
    elif cache.exists():
        text = cache.read_text(encoding="utf-8")
        log(f"図面PDF: 保存済みの OCR 結果（{cache.name}）を使用")
    if text is None:
        try:
            from pypdf import PdfReader
            text = "".join((pg.extract_text() or "") for pg in PdfReader(str(pdf_path)).pages)
            log("図面PDF: OCR が使えないため PDF 内の文字情報で照合")
        except Exception as e:
            log(f"図面PDFを読めないため符号照合をスキップ（{e}）")
            return out
    fig_syms = set(re.findall(r"(?<![0-9A-Za-z.])(\d{1,4}[A-Z]?)(?![0-9A-Za-z.])", text.translate(Z2H)))
    if not fig_syms:
        log("図面PDFに文字情報がないため符号照合をスキップ")
        return out
    if table is None:
        table = symbol_table(paras)
    # 対応表と同じ抽出（文献番号・請求項などは除く）を使う
    spec_syms = {num: (e["occ"][0][0], e["fw"], e["occ"][0][2], e["occ"][0][3]) for num, e in table.items() if e["occ"]}
    # 明細書→図面の向きで照合（図面側の数字は目盛りなども含むため逆向きは使わない）
    missing = [(k, v) for k, v in spec_syms.items() if k not in fig_syms
               and "図示しない" not in v[2].text[max(0, v[3] - 15):v[3]]]
    if spec_syms and len(missing) > len(spec_syms) / 2:
        log(f"明細書の符号 {len(spec_syms)} 個のうち {len(missing)} 個が図面の文字に見つからず、図面PDFの文字が読めていないとみなして照合を省略")
        return out
    for num in spec_syms:
        table[num]["fig"] = num in fig_syms
    for num, (term, fw, p, i) in missing:
        table[num]["issues"].append("図面にない")
        out.append(Finding("rule", p.label, term + fw, "", "図面・符号との一致",
                           f"符号「{term}{fw}」が図面（PDFの文字）に見つからない。 要確認", p.idx, i, needs_check=True))
    log(f"図面PDFの符号 {len(fig_syms)} 個と照合")
    return out

# ---------------------------------------------------------------------------
# 3. LLM
# ---------------------------------------------------------------------------
# LLM を呼ぶ直前に呼ばれる関数（アプリが「PaddleOCR が止まるまで待つ」「キャンセル」に使う）
BEFORE_CALL = None


class LLM:
    def __init__(self, base_url: str, model: str, api_key: str | None, verbose: bool):
        self.url = base_url.rstrip("/") + "/chat/completions"
        self.model, self.api_key, self.verbose = model, api_key, verbose
        self.calls = 0
        self.stats: list[dict] = []  # 呼び出しごとの所要時間・トークン数（ボトルネックの確認用）

    def short(self, since: int = 0) -> str:
        st = self.stats[since:]
        return f"{len(st)} 回 ・ 平均 {sum(x['secs'] for x in st) / len(st):.1f} 秒" if st else "呼び出しなし"

    def summary(self, since: int = 0) -> str:
        st = self.stats[since:]
        if not st:
            return "LLM 呼び出しなし"
        secs = sum(x["secs"] for x in st)
        pt = sum(x["prompt"] for x in st)
        ct = sum(x["completion"] for x in st)
        cut = sum(1 for x in st if x["finish"] == "length")
        msg = f"{len(st)} 回 / 平均 {secs / len(st):.1f} 秒 / 入力 {pt:,} トークン / 出力 {ct:,} トークン"
        if ct and secs:
            msg += f"（出力 {ct / secs:.1f} トークン/秒・1本あたり）"
        if cut:
            msg += f" / 出力上限で切れた応答 {cut} 回"
        return msg

    def chat(self, system: str, user: str, max_tokens: int = 4096, reasoning: str | None = None,
             temperature: float = 0.1, top_p: float = 0.3) -> str:
        if BEFORE_CALL:
            BEFORE_CALL()
        effort = REASONING if reasoning is None else reasoning
        body = json.dumps({
            "model": self.model,
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
            "temperature": temperature, "top_p": top_p, "max_tokens": max_tokens, "stream": False,
            **({"reasoning_effort": effort} if effort else {}),
        }).encode()
        headers = {"Content-Type": "application/json"}
        if self.api_key:
            headers["Authorization"] = "Bearer " + self.api_key
        delay = 5
        for attempt in range(8):
            try:
                t0 = time.time()
                req = urllib.request.Request(self.url, data=body, headers=headers)
                with urllib.request.urlopen(req, timeout=1800) as r:
                    data = json.loads(r.read())
                self.calls += 1
                choice = data["choices"][0]
                usage = data.get("usage") or {}
                self.stats.append({"secs": time.time() - t0, "prompt": usage.get("prompt_tokens", 0),
                                   "completion": usage.get("completion_tokens", 0), "finish": choice.get("finish_reason", "")})
                text = re.sub(r"<think>.*?</think>", "", choice["message"]["content"] or "", flags=re.S).strip()
                if choice.get("finish_reason") == "length" and not text:
                    log(f"  LLM: 推論だけで出力上限（{max_tokens} トークン）に達し、回答が空でした")
                return text
            except urllib.error.HTTPError as e:
                msg = e.read().decode(errors="replace")[:300]
                if e.code in (429, 500, 502, 503, 504):
                    log(f"  LLM {e.code}（混雑など）: {delay}秒後に再試行 … {msg}")
                    time.sleep(delay)
                    delay = min(delay * 2, 120)
                    continue
                raise SystemExit(f"LLM エラー {e.code}: {msg}")
            except urllib.error.URLError as e:
                if attempt == 0:
                    raise SystemExit(f"LLM に接続できません（{self.url}）: {e.reason}\n"
                                     "  --base-url を確認するか、--no-llm でルールチェックだけ実行してください。")
                time.sleep(delay)
        raise SystemExit("LLM の再試行回数を超えました。")


LINE_FORMAT = "段落番号｜修正前（原文から該当箇所を含む10〜30字をそのまま引用）｜修正後｜種類（タイポ／不自然な表現／用語の不統一／図面・符号との一致）｜備考（迷う場合は「要確認」と理由）"

def _sys_view(name: str, what: str, scope: str) -> str:
    return f"""あなたは特許明細書の事務校正担当者です。下の校正方針に従い、{scope}について「{name}」の観点だけを確認してください。

## 校正方針
{MANUAL}

## 今回確認する観点: {name}
{what}

## 出力
1行に1件、次の形式だけを書く（表や前置きは書かない）。該当がなければ「なし」。
{LINE_FORMAT}"""


# 観点ごとの確認（文書全体の整合を見る観点は全文、局所的な観点は区切りごと）
VIEWS_WHOLE = [
    ("用語の不統一", "同じものを指す名詞（構成要素の名称など）が、箇所によって異なる名称・表記になっていないか。"),
    ("表記ゆれ", "同じ語の漢字／ひらがな／カタカナの書き方が文書内で揺れていないか。"),
    ("符号と名称の不一致", "符号と名称の対応が文書内で食い違っていないか（同じ符号に別の名称、同じ構成要素に別の符号など）。図番号の記載に食い違いがないか。"),
]
VIEWS_LOCAL = [
    ("誤字・誤変換・助詞の誤り", "誤変換、誤字・脱字、助詞（てにをは）の誤りがないか。"),
    ("口語的・明らかに不自然な表現", "口語的な表現や、日本語として明らかに不自然で説明できる箇所がないか。"),
]

def sys_verify() -> str:
    return f"""あなたは特許明細書の事務校正の2次チェック担当です。1次チェックの候補を1件ずつ検証し、採用するか判定してください。

## 校正方針
{MANUAL}

## 判定基準
- 採用: 原文が実際にその通り書かれており、校正方針の確認項目に当たる誤りだと説明できる。
- 不採用: 原文に該当箇所がない、原文が正しい、好みの言い換えにすぎない、技術内容の判断が必要、校正方針で確認しない項目に当たる。

## 出力
候補ごとに1行、次の形式だけを書く。
候補番号｜採用 または 不採用｜種類｜修正後（迷う場合は「要確認」と理由）"""


def parse_lines(text: str) -> list[list[str]]:
    rows = []
    for line in text.splitlines():
        line = line.strip().strip("|").strip()
        if "｜" not in line and "|" not in line:
            continue
        cols = [c.strip() for c in re.split(r"[｜|]", line)]
        if cols and cols[0] in ("段落番号", "候補番号", "No", ":---", "---"):
            continue
        if all(set(c) <= set(":-") for c in cols):
            continue
        rows.append(cols)
    return rows


def rows_to_findings(rows, source: str) -> list[Finding]:
    out = []
    for c in rows:
        if len(c) >= 5 and LABEL_RE.match(c[0]):
            label, before, after, kind, note = c[0], c[1], c[2], c[3], "｜".join(c[4:])
        elif len(c) >= 4 and LABEL_RE.match(c[0]):
            label, before, after, kind, note = c[0], c[1], c[2], c[3], ""
        elif len(c) >= 3:  # 段落番号が抜けた行
            label, before, after, kind, note = "", c[0], c[1], c[2] if len(c) > 2 else "", "｜".join(c[3:])
        else:
            continue
        if re.fullmatch(rf"【?[{FW_DIGIT}0-9]+】?|【[^】]*】", re.sub(r"\s", "", before)):
            continue  # 段落番号だけの引用は位置を特定できないので捨てる
        if kind not in TYPES:
            note = (kind + " " + note).strip()
            kind = "タイポ"
        out.append(Finding(source, label, before, after, kind, note, needs_check="要確認" in note))
    return out


def _verify(llm: LLM, text: str, cands: list[Finding]) -> list[Finding]:
    if not cands:
        return []
    listing = "\n".join(f"{i}｜{c.label}｜{c.before}｜{c.after}｜{c.note}" for i, c in enumerate(cands, 1))
    ver = parse_lines(llm.chat(sys_verify(), f"## 本文\n{text}\n\n## 1次チェックの候補\n{listing}"))
    verdict = {}
    for c in ver:
        num = re.sub(r"\D", "", c[0].translate(str.maketrans("０１２３４５６７８９", "0123456789")))
        if num:
            verdict[int(num)] = c
    out = []
    for i, cand in enumerate(cands, 1):
        v = verdict.get(i)
        if not v or len(v) < 2 or "不採用" in v[1] or "採用" not in v[1]:
            continue
        if len(v) >= 3 and v[2] in TYPES:
            cand.kind = v[2]
        if len(v) >= 4 and "要確認" in v[3]:
            cand.needs_check = True
            cand.note = (cand.note + " " + v[3]).strip()
        out.append(cand)
    return out


VOTE_RUNS = int(os.environ.get("PROOF_VOTE_RUNS", "3"))        # 同じ観点を何回実行するか
PARALLEL = int(os.environ.get("PROOF_LLM_PARALLEL", "2"))        # LLM の同時実行数（ルーターの上限に合わせる）
SAMPLE_TEMP = float(os.environ.get("PROOF_SAMPLE_TEMPERATURE", "0.7"))
# 全文を渡す観点（VIEWS_WHOLE）を使うか。2万字級の明細書では1回1.5〜2分かかるのに指摘が0件だったため既定は off。
# 用語の不統一・符号と名称・表記ゆれはルールチェック（対応表・形態素解析・カタカナ表記ゆれ）で機械的に見る
WHOLE_VIEWS = os.environ.get("PROOF_WHOLE_VIEWS", "0") == "1"


def segments(paras: list[Para], seg_chars: int) -> list[list[Para]]:
    groups, cur, n = [], [], 0
    for p in paras:
        if n + len(p.text) > seg_chars and cur:
            groups.append(cur)
            cur, n = [], 0
        cur.append(p)
        n += len(p.text)
    if cur:
        groups.append(cur)
    return groups


def llm_vote(llm: LLM, paras: list[Para], seg_chars: int):
    """観点別チェックを VOTE_RUNS 回ずつ（サンプリングを変えて）実行し、検出箇所ごとに票を数える。
    判定規則（事前に固定）: 2回以上検出 → 採用 / 1回だけ → 2次チェック（llm_verify）で採用されたものだけ採用。
    戻り値: (採用, {区切り番号: 1票の候補}, 区切りの本文)"""
    from concurrent.futures import ThreadPoolExecutor

    full = "\n".join(p.text for p in paras)
    groups = segments(paras, seg_chars)
    group_of = {p.idx: gi for gi, g in enumerate(groups) for p in g}
    gtext = ["\n".join(p.text for p in g) for g in groups]

    # 同じプロンプト（観点 → 区切り → 回）を続けて投げる。LLM サーバーのプロンプトキャッシュが効き、
    # 2回目以降は本文の読み込み（プロンプト処理）を省ける。各回は独立にサンプリングするので結果の性質は変わらない
    tasks = []  # (run, title, system, user, source)
    if WHOLE_VIEWS:
        for name, what in VIEWS_WHOLE:
            for run in range(VOTE_RUNS):
                tasks.append((run, f"全文「{name}」", _sys_view(name, what, "明細書全体"), full, "llm-whole"))
    for name, what in VIEWS_LOCAL:
        system = _sys_view(name, what, "渡された本文の一部")
        for gi in range(len(groups)):
            for run in range(VOTE_RUNS):
                tasks.append((run, f"区切り{gi + 1}「{name}」", system, "## 本文\n" + gtext[gi], "llm-seg"))
    done = [0]

    def run_task(t):
        run, title, system, user, source = t
        out = rows_to_findings(parse_lines(llm.chat(system, user, temperature=SAMPLE_TEMP, top_p=0.95)), source)
        done[0] += 1
        log(f"LLM: {done[0]}/{len(tasks)} {run + 1}回目 {title} … {len(out)}件")
        for f in out:
            f.run = run
        return out

    log(f"LLM: 観点別チェックを {VOTE_RUNS} 回ずつ（{len(tasks)} 呼び出し、同時 {PARALLEL}）" + ("" if WHOLE_VIEWS else "。全文の観点はルールチェックで代替"))
    cands: list[Finding] = []
    with ThreadPoolExecutor(PARALLEL) as ex:
        for out in ex.map(run_task, tasks):
            cands += out

    # 原文の位置を特定し、同じ段落で範囲が重なる候補を1つにまとめて得票を数える
    located = [f for f in cands if locate(f, paras)]
    clusters: list[list[Finding]] = []
    for f in sorted(located, key=lambda f: (f.para_idx, f.start or 0)):
        fs, fe = f.start or 0, (f.start or 0) + max(len(f.before), 1)
        for c in clusters:
            g = c[0]
            gs, ge = g.start or 0, (g.start or 0) + max(len(g.before), 1)
            if g.para_idx == f.para_idx and fs < ge and gs < fe:
                c.append(f)
                break
        else:
            clusters.append([f])

    adopted: list[Finding] = []
    single: dict[int, list[Finding]] = {}
    for c in clusters:
        votes = len({getattr(f, "run", 0) for f in c})
        afters = {}
        for f in c:
            afters[norm(f.after)] = afters.get(norm(f.after), 0) + 1
        best_after = max(afters, key=afters.get)
        rep = next(f for f in c if norm(f.after) == best_after)
        rep.votes = votes
        rep.note = (rep.note + f" （{VOTE_RUNS}回中{votes}回検出）").strip()
        if votes >= 2:
            adopted.append(rep)
        else:
            single.setdefault(group_of.get(rep.para_idx, 0), []).append(rep)
    log(f"LLM: 候補 {len(clusters)} 箇所（2回以上 {len(adopted)}、1回だけ {sum(len(v) for v in single.values())}）")
    return adopted, single, gtext


def llm_verify(llm: LLM, single: dict[int, list[Finding]], gtext: list[str]) -> list[Finding]:
    from concurrent.futures import ThreadPoolExecutor

    def verify_group(item):
        gi, lst = item
        return _verify(llm, gtext[gi], lst)

    out: list[Finding] = []
    with ThreadPoolExecutor(PARALLEL) as ex:
        for got in ex.map(verify_group, sorted(single.items())):
            out += got
    log(f"LLM: 1票の候補 {sum(len(v) for v in single.values())} 件のうち {len(out)} 件を検証で採用")
    return out


def llm_checks(llm: LLM, paras: list[Para], seg_chars: int) -> list[Finding]:
    adopted, single, gtext = llm_vote(llm, paras, seg_chars)
    adopted += llm_verify(llm, single, gtext)
    log(f"LLM: 多数決＋検証の結果 {len(adopted)} 件")
    return adopted


# ---------------------------------------------------------------------------
# 4. 位置特定・重複除去
# ---------------------------------------------------------------------------
def norm(s: str) -> str:
    return re.sub(r"[\s　]", "", s)


def locate(f: Finding, paras: list[Para]) -> bool:
    if f.para_idx is not None:
        return True
    b = norm(f.before)
    if len(b) < 2:
        return False
    cands = sorted(paras, key=lambda p: 0 if p.label == f.label else 1)
    for p in cands:
        i = p.text.find(f.before)
        if i < 0:
            i = norm(p.text).find(b)
            if i >= 0:
                i = p.text.find(f.before.strip()[:4]) if p.text.find(f.before.strip()[:4]) >= 0 else 0
        if i >= 0 and b in norm(p.text):
            f.para_idx, f.start, f.label = p.idx, i, p.label  # 段落番号は原文側の位置から決め直す
            return True
    return False  # 原文に存在しない引用（LLM の幻覚）は捨てる


def dedupe(findings: list[Finding]) -> list[Finding]:
    order = {"rule": 0, "llm-whole": 1, "llm-seg": 2}
    findings.sort(key=lambda f: order.get(f.source, 9))
    kept: list[Finding] = []
    def span(x):
        s = x.start or 0
        return s, s + max(len(x.before), 1)

    for f in findings:
        k = norm(f.before)
        dup = False
        fs, fe = span(f)
        for g in kept:
            if g.para_idx != f.para_idx:
                continue
            kg = norm(g.before)
            gs, ge = span(g)
            # 同じ段落で引用範囲が重なれば同じ指摘とみなす（二重の変更履歴を防ぐ）
            if k in kg or kg in k or (fs < ge and gs < fe) or (f.after and norm(f.after) == norm(g.after)):
                dup = True
                break
        if not dup:
            kept.append(f)
    kept.sort(key=lambda f: (f.para_idx or 0, f.start or 0))
    return kept


# ---------------------------------------------------------------------------
# 5. docx へ反映（変更履歴 / コメント）
# ---------------------------------------------------------------------------
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def qn(tag: str) -> str:
    return "{%s}%s" % (W_NS, tag)


def simple_run(r) -> bool:
    kids = [c for c in r._r if c.tag != qn("rPr")]
    return len(kids) == 1 and kids[0].tag == qn("t")


def isolate(paragraph, start: int, end: int):
    """paragraph.text[start:end] をちょうど覆う run 列を作って返す（できなければ None）。"""
    pos = 0
    runs = list(paragraph.runs)
    if "".join(r.text for r in runs) != paragraph.text:
        return None  # ハイパーリンク等を含む段落は扱わない
    out = []
    for r in runs:
        t = r.text
        a, b = pos, pos + len(t)
        pos = b
        if b <= start or a >= end or not t:
            continue
        if not simple_run(r):
            return None
        s, e = max(start, a) - a, min(end, b) - a
        if s > 0:  # 前半を切り離す
            head = copy.deepcopy(r._r)
            r._r.addprevious(head)
            head.find(qn("t")).text = t[:s]
            head.find(qn("t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if e < len(t):  # 後半を切り離す
            tail = copy.deepcopy(r._r)
            r._r.addnext(tail)
            tail.find(qn("t")).text = t[e:]
            tail.find(qn("t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r._r.find(qn("t")).text = t[s:e]
        r._r.find(qn("t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        out.append(r)
    return out or None


_rev_id = [9000]


def tracked_replace(paragraph, start: int, end: int, new: str, date: str) -> bool:
    runs = isolate(paragraph, start, end) if end > start else None
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn as dqn
    if end > start and not runs:
        return False
    if end == start:  # 挿入のみ: start 位置で run を分割
        runs_after = isolate(paragraph, start, start + 1) if start < len(paragraph.text) else None
        if not runs_after:
            return False
        anchor, before = runs_after[0]._r, True
        tmpl = anchor
    else:
        anchor, before = runs[-1]._r, False
        tmpl = runs[0]._r
        dele = OxmlElement("w:del")
        _rev_id[0] += 1
        dele.set(dqn("w:id"), str(_rev_id[0]))
        dele.set(dqn("w:author"), AUTHOR)
        dele.set(dqn("w:date"), date)
        runs[0]._r.addprevious(dele)
        for r in runs:
            t = r._r.find(qn("t"))
            t.tag = qn("delText")
            dele.append(r._r)
        anchor = dele
    if new:
        ins = OxmlElement("w:ins")
        _rev_id[0] += 1
        ins.set(dqn("w:id"), str(_rev_id[0]))
        ins.set(dqn("w:author"), AUTHOR)
        ins.set(dqn("w:date"), date)
        nr = copy.deepcopy(tmpl)
        for c in list(nr):
            if c.tag != qn("rPr"):
                nr.remove(c)
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = new
        nr.append(t)
        ins.append(nr)
        if before:
            anchor.addprevious(ins)
        else:
            anchor.addnext(ins)
    return True


def minimal_edit(before: str, after: str) -> tuple[int, int, str]:
    i = 0
    while i < min(len(before), len(after)) and before[i] == after[i]:
        i += 1
    j = 0
    while j < min(len(before), len(after)) - i and before[-1 - j] == after[-1 - j]:
        j += 1
    return i, len(before) - j, after[i:len(after) - j]


def apply_to_docx(doc, findings: list[Finding]):
    date = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    # 段落ごとに後ろから適用して位置ずれを防ぐ
    for f in sorted(findings, key=lambda f: (f.para_idx, -(f.start or 0))):
        p = doc.paragraphs[f.para_idx]
        text = p.text
        bi = text.find(f.before)
        if bi < 0:
            bi = f.start or 0
        comment = f"【{f.kind}】{f.before} → {f.after or '（要確認）'}"
        if f.votes and f.votes < VOTE_RUNS:
            comment = "【要確認】" + comment
        if f.note:
            comment += f"\n{f.note}"
        done = False
        unanimous = f.votes == 0 or f.votes >= VOTE_RUNS  # 全回一致の指摘とルールの指摘だけ変更履歴の対象
        if f.after and not f.needs_check and unanimous and text.find(f.before) >= 0:
            s, e, new = minimal_edit(f.before, f.after)
            old = f.before[s:e]
            kanji_swap = len(old) >= 2 and len(new) >= 2 and re.fullmatch(r"[一-龠々]+", old) and re.fullmatch(r"[一-龠々]+", new)
            small = (e - s) <= 3 and len(new) <= 3 or (e == s and len(new) <= 6)  # 数文字の差し替え・短い補完だけ
            if small and not kanji_swap:  # 小さな修正だけ変更履歴にする（意味の変わる語の入れ替えはコメント）
                done = tracked_replace(p, bi + s, bi + e, new, date)
                if done:
                    f.applied = "change"
        if not done:
            end = bi + len(f.before) if text.find(f.before) >= 0 else min(len(text), bi + 10)
            runs = isolate(p, bi, end) or [r for r in p.runs if r.text][:1]
            if runs:
                doc.add_comment(runs, text=comment, author=AUTHOR, initials="AI")
                f.applied = "comment"
            else:
                f.applied = "report-only"


def write_table_csv(path, table: dict[str, dict]):
    with open(path, "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["符号", "本文の名称（回数）", "符号の説明", "初出", "図面", "矛盾"])
        w.writerows(table_rows(table))


# ---------------------------------------------------------------------------
def log(msg: str):
    print(msg, file=sys.stderr, flush=True)


def main():
    ap = argparse.ArgumentParser(description="特許明細書の事務校正（ルール + ローカルLLM）")
    ap.add_argument("docx", type=Path)
    ap.add_argument("--base-url", default=DEFAULT_BASE_URL)
    ap.add_argument("--model", default=DEFAULT_MODEL)
    ap.add_argument("--api-key", default=os.environ.get("PROOF_LLM_API_KEY"))
    ap.add_argument("--no-llm", action="store_true", help="ルールチェックだけ実行する")
    ap.add_argument("--segment-chars", type=int, default=1200)
    ap.add_argument("--out-dir", type=Path)
    args = ap.parse_args()

    from docx import Document

    src = args.docx
    doc = Document(str(src))
    paras = extract_paragraphs(doc)
    if not paras:
        raise SystemExit("【書類名】明細書 以降の本文が見つかりません。")
    log(f"本文 {len(paras)} 段落 / {sum(len(p.text) for p in paras)} 文字")

    table = symbol_table(paras)
    findings = rule_checks(paras, table)
    log(f"ルールチェック: {len(findings)} 件")
    if not args.no_llm:
        llm = LLM(args.base_url, args.model, args.api_key, verbose=True)
        t0 = time.time()
        findings += llm_checks(llm, paras, args.segment_chars)
        log(f"LLM {llm.summary()} / 合計 {int(time.time() - t0)} 秒")

    located = [f for f in findings if locate(f, paras)]
    dropped = len(findings) - len(located)
    final = dedupe(located)
    log(f"指摘 {len(final)} 件（原文に見つからない引用 {dropped} 件を除外）")

    apply_to_docx(doc, final)
    out_dir = args.out_dir or src.parent
    out_docx = out_dir / f"【AI校正】{src.name}"
    out_csv = out_dir / f"{src.stem}_修正箇所一覧.csv"
    doc.save(str(out_docx))
    write_table_csv(out_dir / f"{src.stem}_符号対応表.csv", table)
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["No", "段落番号", "修正前", "修正後", "種類", "備考", "要確認", "反映", "検出元"])
        for i, f in enumerate(final, 1):
            w.writerow([i, f.label, f.before, f.after, f.kind, f.note, "○" if f.needs_check else "", f.applied, f.source])
    print(out_docx)
    print(out_csv)
    print(out_dir / f"{src.stem}_符号対応表.csv")


if __name__ == "__main__":
    main()
